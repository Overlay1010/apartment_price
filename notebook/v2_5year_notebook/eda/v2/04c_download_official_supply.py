"""
국토교통부 주택건설 인허가·준공실적 공식 통계 수집 및 서울 값 추출 (통계누리)

목적: 주택인허가 원장(구×월 Feature)을 공식 통계와 대조하기 위한 서울 기준값 확보.

출처 (statMetaView.do 첨부 목록)
- 인허가: hRsId=31,  hFormId=1948
- 준공  : hRsId=468, hFormId=5373
  (착공 hRsId=471/hFormId=5387 도 같은 구조로 올라와 있지만 여기서는 받지 않는다.)

받는 파일 (data/external/supply_official/, 있으면 다시 받지 않음)
- molit_{kind}_zip_YYYY.zip      : 2019~2025 연도별 zip (그 해 1~12월 월별 xls 12개)
- molit_{kind}_YYYY12.xls        : 위 zip에서 꺼낸 12월(12월말 누계) 파일
- molit_{kind}_2026MM.xls        : 2026년 월별 파일 (첨부 목록에 있는 달 전부, 최신월 포함)
- molit_{kind}_ts_*.xlsx         : 연도별 시계열 파일(확정치, '(90년-25년)…' 등 최신판)
  kind = permit(인허가) | completion(준공)

추출 결과
- molit_seoul_permit_completion.csv
  kind, year, month(연간은 빈칸), total, apartment, source_file, basis
  basis:
    dec_file_ytd           각 연도 12월 파일 '참고1 주택유형별 대분류' 12월말 누계 (당시 공표치)
    latest_file_ytd        2026 최신월 파일의 연초~해당월 누계 (month 칸 = 누계 마지막 달)
    timeseries_type        시계열 '연도별,유형별' 파일의 서울 합계/아파트 (최신 확정치)
    monthly_file           각 월 파일 '참고1' 월계 블록 (그 달 최초 공표치, 아파트 포함)
    timeseries_month_region 시계열 '연도별,월별,지역별' 파일의 서울 월별 (확정치, 아파트 없음)
  같은 (kind, year, month) 라도 파일·시점마다 수정(확정치 반영)될 수 있으므로 basis/source_file로 구분한다.
- _manifest.csv : file_name, source_url, ref

zip 내부 파일명은 cp437로 깨져 있어 cp949로 다시 디코딩한다 (04b_download_unsold.py와 동일).

실행: python notebook/v2_5year_notebook/eda/v2/04c_download_official_supply.py
"""

import csv
import io
import re
import time
import zipfile
from pathlib import Path
from urllib.parse import quote, urlencode

import openpyxl
import requests
import xlrd


def find_project_root(start=Path(__file__).resolve()):
    for path in [start, *start.parents]:
        if (path / "data").exists() and (path / ".env").exists():
            return path
    raise FileNotFoundError("프로젝트 루트를 찾을 수 없습니다.")


PROJECT_DIR = find_project_root()
OUT_DIR = PROJECT_DIR / "data" / "external" / "supply_official"
RESULT_PATH = OUT_DIR / "molit_seoul_permit_completion.csv"
MANIFEST_PATH = OUT_DIR / "_manifest.csv"

BASE = "https://stat.molit.go.kr"
DOWNLOAD_URL = f"{BASE}/portal/common/downLoadFile.do"
UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/120.0 Safari/537.36"
)
DELAY = 2.0  # 요청 사이 지연(초)

YEARS = range(2019, 2026)
LATEST_YEAR = 2026

KINDS = {
    "permit": {
        "label": "인허가",
        "meta": f"{BASE}/portal/cate/statMetaView.do?hRsId=31&hFormId=1948",
        # 시계열 파일: (첨부명 정규식, 저장명)
        "ts": [
            (r"^\(90년-(\d\d)년\)인허가_연도별,월별,지역별\.xlsx$", "molit_permit_ts_month_region_1990_20{}.xlsx"),
            (r"^\(90년-(\d\d)년\)인허가_연도별,지역별\.xlsx$", "molit_permit_ts_region_1990_20{}.xlsx"),
            (r"^\(95년-(\d\d)년\)인허가_연도별,유형별\.xlsx$", "molit_permit_ts_type_1995_20{}.xlsx"),
        ],
    },
    "completion": {
        "label": "준공",
        "meta": f"{BASE}/portal/cate/statMetaView.do?hRsId=468&hFormId=5373",
        "ts": [
            (r"^\(05년-(\d\d)년\)준공_연도별,월별,지역별\.xlsx$", "molit_completion_ts_month_region_2005_20{}.xlsx"),
            (r"^\(05년-(\d\d)년\)준공_연도별,지역별\.xlsx$", "molit_completion_ts_region_2005_20{}.xlsx"),
            (r"^\(05년-(\d\d)년\)준공_연도별,유형별\.xlsx$", "molit_completion_ts_type_2005_20{}.xlsx"),
        ],
    },
}

session = requests.Session()
session.headers.update({"User-Agent": UA})
manifest = {}  # file_name -> (source_url, ref)


# ---------------------------------------------------------------- 다운로드
def download_url(file_name):
    query = urlencode(
        {"oFileName": file_name, "rFileName": file_name, "midpath": "/stat_file/"},
        quote_via=quote,
    )
    return f"{DOWNLOAD_URL}?{query}"


def fetch(file_name, referer):
    """첨부 파일을 받는다. 없는 파일이면 None (서버가 HTML 안내 페이지를 돌려준다)."""
    time.sleep(DELAY)
    resp = session.get(download_url(file_name), headers={"Referer": referer}, timeout=300)
    resp.raise_for_status()
    # xls(OLE2: D0CF11E0) / xlsx·zip(PK) 시그니처 확인
    if not (resp.content.startswith(b"PK") or resp.content.startswith(b"\xd0\xcf\x11\xe0")):
        return None
    return resp.content


def list_attachments(meta_url):
    resp = session.get(meta_url, timeout=60)
    resp.raise_for_status()
    return re.findall(r"downFile\('([^']+)'", resp.text)


def save_attachment(name, out, referer, ref):
    if not out.exists():
        content = fetch(name, referer)
        if content is None:
            print(f"[누락] {name} 다운로드 실패")
            return False
        out.write_bytes(content)
        print(f"받음: {name} -> {out.name}")
    manifest[out.name] = (download_url(name), ref)
    return True


def decode_member_name(info):
    if info.flag_bits & 0x800:  # UTF-8 플래그가 있으면 그대로
        return info.filename
    try:
        return info.filename.encode("cp437").decode("cp949")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return info.filename


def month_of_member(name):
    """'24년 03월 인허가실적(홈페이지게시용).xls' / '20년_12월_…' -> (2024, 3)"""
    m = re.search(r"(\d{2})년[\s_]*(\d{1,2})월", name)
    return (2000 + int(m.group(1)), int(m.group(2))) if m else None


def zip_members(zip_path):
    """{(y, m): (member_name, ZipInfo)} — 같은 달이 여러 개면 '(수정)' 파일 우선."""
    out = {}
    with zipfile.ZipFile(zip_path) as zf:
        for info in zf.infolist():
            name = decode_member_name(info)
            ym = month_of_member(name)
            if ym is None or not name.lower().endswith((".xls", ".xlsx")):
                continue
            if ym not in out or "(수정)" in name:
                out[ym] = (name, info)
    return out


def download_all():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for kind, cfg in KINDS.items():
        meta = cfg["meta"]
        attachments = list_attachments(meta)
        print(f"[{cfg['label']}] 첨부 {len(attachments)}건")

        # 1) 2019~2025 연도별 zip + 12월 파일
        for year in YEARS:
            zip_name = next(
                (a for a in attachments if a.startswith(str(year)) and a.endswith(".zip")), None
            )
            if zip_name is None:
                print(f"[누락] {cfg['label']} {year}년 zip이 첨부 목록에 없음")
                continue
            zip_path = OUT_DIR / f"molit_{kind}_zip_{year}.zip"
            if not save_attachment(zip_name, zip_path, meta, f"{year}년 1~12월 월별 파일 묶음"):
                continue
            members = zip_members(zip_path)
            if (year, 12) not in members:
                print(f"[누락] {zip_path.name} 안에 12월 파일 없음")
                continue
            name, info = members[(year, 12)]
            out = OUT_DIR / f"molit_{kind}_{year}12{Path(name).suffix.lower()}"
            if not out.exists():  # molit_permit_202112.xls 처럼 이미 있으면 덮어쓰지 않음
                with zipfile.ZipFile(zip_path) as zf:
                    out.write_bytes(zf.read(info))
                print(f"추출: {zip_path.name}#{name} -> {out.name}")
            manifest[out.name] = (f"{download_url(zip_name)}#{name}", f"{year}-12 (12월말 누계)")

        # 2) 2026 월별 파일 (첨부 목록에 있는 달 전부)
        yy = LATEST_YEAR % 100
        for a in attachments:
            m = re.match(rf"^{yy}년 (\d\d)월 .*\.(xlsx?)$", a)
            if m:
                out = OUT_DIR / f"molit_{kind}_{LATEST_YEAR}{m.group(1)}.{m.group(2)}"
                save_attachment(a, out, meta, f"{LATEST_YEAR}-{m.group(1)} (월말 누계)")

        # 3) 시계열(확정치) 파일
        for pattern, dst in cfg["ts"]:
            hits = sorted((re.match(pattern, a).group(1), a) for a in attachments if re.match(pattern, a))
            if not hits:
                print(f"[누락] 시계열 파일 패턴 {pattern}")
                continue
            last_yy, name = hits[-1]  # 가장 최근판
            save_attachment(name, OUT_DIR / dst.format(last_yy), meta, f"시계열 ~20{last_yy}년 (확정치)")


# ---------------------------------------------------------------- 추출
def norm(v):
    return re.sub(r"[\s\xa0]+", "", str(v)) if v is not None else ""


def num(v):
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None


def parse_reference_sheet(book):
    """'참고1 주택유형별 대분류' 시트에서 서울 블록별 (계, 아파트) 를 읽는다.

    반환: {"ytd": (label_ym, total, apt), "prev_ytd": (...), "month": (...)}
    시트는 [누계(당월) | 누계(전월) | 월계(당월)] 3개 블록이 가로로 놓여 있다
    (2026-07 인허가 파일처럼 전월 누계 블록이 빠지고 한 칸 밀린 경우도 있다).
    블록 제목의 날짜는 원본에서 틀린 경우가 있어('2020년 01말', 전월 누계에 '2019년 10월말' 등)
    기준월은 파일명에서 정하고, label_ym 은 참고용으로만 둔다.
    """
    sheet = next(s for s in book.sheets() if "대분류" in s.name)
    rows = [sheet.row_values(r) for r in range(sheet.nrows)]
    hdr_r = next(r for r, row in enumerate(rows) if {"구분", "아파트"} <= set(map(norm, row)))
    hdr = [norm(v) for v in rows[hdr_r]]
    starts = [c for c, v in enumerate(hdr) if v == "구분"]
    result = {}
    for i, c0 in enumerate(starts):
        c1 = starts[i + 1] if i + 1 < len(starts) else len(hdr)
        label = "".join(norm(rows[r][c]) for r in range(max(0, hdr_r - 3), hdr_r) for c in range(c0, c1))
        dm = re.search(r"(\d{4})년(\d{1,2})월?말", label)
        ym = (int(dm.group(1)), int(dm.group(2))) if dm else None
        c_tot = next(c for c in range(c0, c1) if hdr[c] == "계")
        c_apt = next(c for c in range(c0, c1) if hdr[c] == "아파트")
        seoul = next(row for row in rows[hdr_r + 1:] if norm(row[c0]) == "서울")
        vals = (ym, num(seoul[c_tot]), num(seoul[c_apt]))
        key = "month" if "월계" in label else ("ytd" if "ytd" not in result else "prev_ytd")
        result[key] = vals
    return result


def check_month(ref, ym, name):
    """같은 파일 안에서 월계 = 당월 누계 - 전월 누계 (1월은 월계 = 누계) 인지 확인."""
    for i in (1, 2):
        ytd, month = ref["ytd"][i], ref["month"][i]
        prev = ref.get("prev_ytd", (None, None, None))[i]
        expect = ytd if ym[1] == 1 else (ytd - prev if prev is not None else month)
        if month != expect:
            print(f"[경고] {name}: 월계 {month} != 누계차 {expect} ({'계' if i == 1 else '아파트'})")
    if ref["month"][0] not in (None, ym):
        print(f"[참고] {name}: 블록 제목 날짜 {ref['month'][0]} (파일 기준월 {ym})")


def extract_rows():
    rows = []
    for kind, cfg in KINDS.items():
        label = cfg["label"]

        # (a) 월별 파일: 2019~2025 zip 안 12개 + 2026 개별 파일 -> 월계(월별) / 12월·최신월 누계(연간)
        for year in YEARS:
            zip_path = OUT_DIR / f"molit_{kind}_zip_{year}.zip"
            if not zip_path.exists():
                continue
            with zipfile.ZipFile(zip_path) as zf:
                for (y, m), (name, info) in sorted(zip_members(zip_path).items()):
                    book = xlrd.open_workbook(file_contents=zf.read(info))
                    ref = parse_reference_sheet(book)
                    check_month(ref, (y, m), name)
                    src = f"{zip_path.name}/{name}"
                    rows.append([label, y, m, ref["month"][1], ref["month"][2], src, "monthly_file"])
                    if m == 12:
                        dec = OUT_DIR / f"molit_{kind}_{year}12{Path(name).suffix.lower()}"
                        rows.append([label, y, "", ref["ytd"][1], ref["ytd"][2], dec.name, "dec_file_ytd"])

        latest = sorted(OUT_DIR.glob(f"molit_{kind}_{LATEST_YEAR}[01][0-9].xls*"))
        for path in latest:
            ref = parse_reference_sheet(xlrd.open_workbook(path))
            y, m = int(path.stem[-6:-2]), int(path.stem[-2:])
            check_month(ref, (y, m), path.name)
            rows.append([label, y, m, ref["month"][1], ref["month"][2], path.name, "monthly_file"])
        if latest:
            ref = parse_reference_sheet(xlrd.open_workbook(latest[-1]))
            y, m = int(latest[-1].stem[-6:-2]), int(latest[-1].stem[-2:])
            rows.append([label, y, m, ref["ytd"][1], ref["ytd"][2], latest[-1].name, "latest_file_ytd"])

        # (b) 시계열 유형별: 연도 시트별 서울 합계/아파트
        for path in OUT_DIR.glob(f"molit_{kind}_ts_type_*.xlsx"):
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                sm = re.match(r"(\d\d)년", ws.title)
                if not sm or 2000 + int(sm.group(1)) not in [*YEARS]:
                    continue
                vals = list(ws.iter_rows(max_row=40, values_only=True))
                hdr = next(r for r in vals if norm(r[0]) == "구분")
                hdr = [norm(v) for v in hdr]
                c_tot = next(c for c, v in enumerate(hdr) if v in ("계", "합계"))
                c_apt = hdr.index("아파트")
                seoul = next(r for r in vals if norm(r[0]) == "서울")
                rows.append([label, 2000 + int(sm.group(1)), "", num(seoul[c_tot]), num(seoul[c_apt]),
                             path.name, "timeseries_type"])

        # (c) 시계열 월별·지역별: 서울 행 월별 전체(아파트 없음)
        for path in OUT_DIR.glob(f"molit_{kind}_ts_month_region_*.xlsx"):
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            for ws in wb.worksheets:
                sm = re.match(r"(\d\d)년", ws.title)
                if not sm or 2000 + int(sm.group(1)) not in [*YEARS]:
                    continue
                vals = list(ws.iter_rows(max_row=60, values_only=True))
                hdr = [norm(v) for v in next(r for r in vals if norm(r[0]) == "구분")]
                seoul = next(r for r in vals if norm(r[0]) == "서울")
                for c, h in enumerate(hdr):
                    mm = re.fullmatch(r"(\d{1,2})월", h)
                    if mm and num(seoul[c]) is not None:
                        rows.append([label, 2000 + int(sm.group(1)), int(mm.group(1)), num(seoul[c]), "",
                                     path.name, "timeseries_month_region"])

    order = {"dec_file_ytd": 0, "latest_file_ytd": 1, "timeseries_type": 2,
             "monthly_file": 3, "timeseries_month_region": 4}
    rows.sort(key=lambda r: (r[0] != "인허가", r[1], r[2] if r[2] != "" else 0, order[r[6]]))
    return rows


def main():
    download_all()
    rows = extract_rows()
    with open(RESULT_PATH, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["kind", "year", "month", "total", "apartment", "source_file", "basis"])
        w.writerows(rows)
    print(f"저장: {RESULT_PATH} ({len(rows)}행)")

    # manifest: 이번 실행에서 확인한 파일 + 폴더에 이미 있던 파일(출처 미상이면 빈칸)
    for p in OUT_DIR.iterdir():
        if p.suffix.lower() in (".xls", ".xlsx", ".zip") and p.name not in manifest:
            manifest[p.name] = ("", "")
    with open(MANIFEST_PATH, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["file_name", "source_url", "ref"])
        for name in sorted(manifest):
            w.writerow([name, *manifest[name]])
    print(f"manifest 저장: {MANIFEST_PATH} ({len(manifest)}건)")

    # 연간 요약 출력 (12월 파일 기준)
    for r in rows:
        if r[6] in ("dec_file_ytd", "latest_file_ytd"):
            print(f"  {r[0]} {r[1]}{'' if r[2] == '' else f'(~{r[2]}월)'}: 전체 {r[3]:,} / 아파트 {r[4]:,}")


if __name__ == "__main__":
    main()
